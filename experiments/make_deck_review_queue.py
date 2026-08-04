"""검수 큐 정리 발표 자료(PPTX) — 그림 한 장이 한 슬라이드.

각 장의 주장은 그림 안의 제목이 이미 말한다. 슬라이드는 그림을 크게 놓고, 아래에
그 주장을 뒷받침하는 수치 한 줄만 붙인다. 글머리표로 설명을 반복하면 발표자가
그림 대신 글을 읽게 된다.

  python experiments/make_deck_review_queue.py [out.pptx]
"""
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

REPO = "/data3/bhkim/workspace/train-ego-path-detection"
FIG = os.path.join(REPO, "figures")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(REPO, "deck_review_queue.pptx")

W, H = Inches(13.333), Inches(7.5)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x8A, 0x8A, 0x8A)
GOOD = RGBColor(0x2E, 0x9E, 0x2E)
BAD = RGBColor(0xE5, 0x39, 0x35)
FONT = "Noto Sans CJK KR"

SLIDES = [
    ("fig_review_composition.png", "① 문제",
     "844프레임 중 719장(85%)이 '어떤 모델이 답을 안 냈다'는 이유였다. "
     "정리 후에는 279장이 실제 모델 불일치다."),
    ("fig_review_models.png", "② 원인",
     "osdar23-seg-r18은 25%, brilliant-horse-15는 15%의 프레임에서 침묵한다. "
     "나머지 다섯은 4% 이하. 두 모델은 답을 냈을 때의 정확도도 가장 나빴다."),
    ("fig_review_frame.png", "③ 근거",
     "왼쪽: 제외한 두 모델만 0.30 빗나가고 남긴 다섯은 정답과 0.0005 이내. "
     "오른쪽: 남긴 다섯이 갈리는 프레임 — 이런 것이 사람에게 간다."),
    ("fig_review_tradeoff.png", "④ 선택",
     "'7개 중 5개만 있으면 통과'는 큐를 230까지 줄이지만 심각한 오류의 21%를 놓친다. "
     "약한 2개 제외는 416으로 줄이면서 검출률 100%를 지킨다."),
    ("fig_review_waterfall.png", "⑤ 내역",
     "−265는 규칙 개선분, −208은 이벤트 9건이 '제외 검토' 대상으로 이관된 몫, "
     "+45는 되살아난 이벤트, +1은 자동 라벨이 없던 프레임."),
    ("fig_review_events.png", "⑥ 이벤트 단위 판단",
     "이벤트 내내 모델이 갈리는 9건은 프레임 검수로 감당할 대상이 아니라서 "
     "학습 목록(55 → 46)에서 제외했다."),
]


def textbox(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    for i, line in enumerate(text.split("\n")):
        r = (p if i == 0 else tf.add_paragraph()).add_run()
        r.text = line
        f = r.font
        f.size, f.bold, f.name = Pt(size), bold, FONT
        f.color.rgb = color
        if i:
            tf.paragraphs[i].alignment = align
    return tb


def put_image(slide, path, top, max_h):
    """가로세로비를 지켜 가운데에 놓는다."""
    with Image.open(path) as im:
        iw, ih = im.size
    avail_w = W - Inches(1.2)
    scale = min(avail_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(path, int((W - w) / 2), top, w, h)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # ---- 표지
    s = blank(prs)
    textbox(s, Inches(0.9), Inches(2.4), W - Inches(1.8), Inches(1.2),
            "검수 큐 정리", 54, INK, bold=True)
    textbox(s, Inches(0.95), Inches(3.5), W - Inches(1.8), Inches(1.0),
            "사람이 봐야 하는 프레임 844 → 417", 28, GOOD, bold=True)
    textbox(s, Inches(0.95), Inches(4.35), W - Inches(1.8), Inches(1.6),
            "앙상블 검수 큐의 85%가 프레임 난이도가 아니라 약한 모델의 침묵으로\n"
            "부풀어 있었다. 근거와 그 결과를 정리한다.", 16, MUTED)
    textbox(s, Inches(0.95), Inches(6.3), W - Inches(1.8), Inches(0.5),
            "TEP-Net · 분기부 자동 라벨 검수", 13, MUTED)

    # ---- 그림 슬라이드
    for name, kicker, caption in SLIDES:
        s = blank(prs)
        textbox(s, Inches(0.6), Inches(0.28), Inches(3.0), Inches(0.45),
                kicker, 15, GOOD, bold=True)
        put_image(s, os.path.join(FIG, name), Inches(0.85), Inches(5.35))
        textbox(s, Inches(0.9), Inches(6.45), W - Inches(1.8), Inches(0.8),
                caption, 14, MUTED)

    # ---- 요약
    s = blank(prs)
    textbox(s, Inches(0.9), Inches(0.7), W - Inches(1.8), Inches(0.8),
            "결과", 40, INK, bold=True)
    rows = [("사람이 볼 프레임", "844", "417", "−51%"),
            ("실제 불일치가 이유인 비율", "15%", "67%", "+52%p"),
            ("심각한 오류(폭 0.10 초과) 검출률", "100%", "100%", "유지"),
            ("학습 이벤트 목록", "55", "46", "불신 9건 제외"),
            ("자동 라벨 없이 방치되던 프레임", "1", "0", "해결")]
    y = Inches(2.15)
    textbox(s, Inches(0.95), Inches(1.72), Inches(6.0), Inches(0.4), "항목", 13, MUTED, bold=True)
    for lab, x in (("이전", 7.3), ("이후", 9.1), ("변화", 10.9)):
        textbox(s, Inches(x), Inches(1.72), Inches(1.6), Inches(0.4), lab, 13, MUTED, bold=True)
    for label, before, after, delta in rows:
        textbox(s, Inches(0.95), y, Inches(6.2), Inches(0.5), label, 16, INK)
        textbox(s, Inches(7.3), y, Inches(1.6), Inches(0.5), before, 16, MUTED)
        textbox(s, Inches(9.1), y, Inches(1.6), Inches(0.5), after, 16, INK, bold=True)
        textbox(s, Inches(10.9), y, Inches(2.0), Inches(0.5), delta, 16, GOOD, bold=True)
        y = Emu(int(y) + int(Inches(0.72)))
    textbox(s, Inches(0.95), Inches(5.9), W - Inches(1.9), Inches(1.0),
            "측정 근거 — 침묵률: 캐시 3981프레임 · 정확도: 손라벨 2268프레임(28 이벤트) · "
            "검출률: 손라벨이 있는 분기 이벤트 6건(486프레임).\n"
            "검출률 차이는 대상 오류가 14~40건이라 한두 건 차이는 잡음 범위로 본다.",
            12, MUTED)

    prs.save(OUT)
    print("wrote", OUT, f"({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
