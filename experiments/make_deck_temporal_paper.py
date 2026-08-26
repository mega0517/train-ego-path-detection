"""논문 발표용 16장 PPT — 학습형 시간 정련 대 이동평균 (KRRI 서식).

서식은 paper_summary_5p.pptx에서 실측해 옮겼다: 상단 #1B262C 밴드에 흰 로고판,
본문은 #F2F5F7 패널에 절 제목만 색으로 구분하고, 강조할 결론만 어두운 패널로
반전시킨다. 맑은 고딕 한 종류로 크기만 달리한다.

수치는 paper_KCI_temporal_2026.md의 표 1~4에서 옮긴 것이며, 표는 텍스트 상자가
아니라 실제 표 개체로 넣어 발표 중 수정이 가능하게 했다.

  python experiments/make_deck_temporal_paper.py [out.pptx]
"""
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

REPO = "/data3/bhkim/workspace/train-ego-path-detection"
FIG = os.path.join(REPO, "figures")
LOGO = os.path.join(REPO, "static", "krri_logo.png")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(REPO, "deck_temporal_paper.pptx")

W, H = Inches(13.333), Inches(7.5)
FONT = "맑은 고딕"

INK = RGBColor(0x1B, 0x26, 0x2C)        # 본문·헤더 밴드
SUB = RGBColor(0x47, 0x50, 0x5C)        # 보조 텍스트
PANEL = RGBColor(0xF2, 0xF5, 0xF7)      # 밝은 패널
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
META = RGBColor(0xB0, 0xBE, 0xC5)       # 헤더 우측 메타
RED = RGBColor(0xC6, 0x28, 0x28)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
TEAL = RGBColor(0x00, 0x83, 0x8F)
LIME = RGBColor(0x9C, 0xCC, 0x65)       # 어두운 패널 위 강조
SKY = RGBColor(0x4F, 0xC3, 0xF7)

HEAD_H = Inches(1.28)
BODY_TOP = Inches(1.50)
MARGIN = Inches(0.30)
CONTENT_W = W - 2 * MARGIN


# ---------------------------------------------------------------- 기본 도형
def rect(slide, x, y, w, h, fill):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def text(slide, x, y, w, h, runs, size=12.5, color=INK, bold=False,
         align=PP_ALIGN.LEFT, spacing=1.0, anchor=MSO_ANCHOR.TOP):
    """runs: 문자열, 또는 문단마다 [(글자, 굵게, 색), ...] 목록."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if i:
            p.space_before = Pt(4)
        for piece in ([(para, bold, color)] if isinstance(para, str) else para):
            s, b, c = piece
            r = p.add_run()
            r.text = s
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = b
            r.font.color.rgb = c
    return tb


def header(slide, title, meta="", page=None):
    rect(slide, 0, 0, W, HEAD_H, INK)
    with Image.open(LOGO) as im:
        lw, lh = im.size
    logo_w, logo_h = Inches(2.60), Inches(2.60) * lh / lw
    rect(slide, Inches(0.25), Inches(0.20), logo_w + Inches(0.14),
         logo_h + Inches(0.14), WHITE)
    slide.shapes.add_picture(LOGO, Inches(0.32), Inches(0.27), logo_w, logo_h)
    text(slide, MARGIN, Inches(0.68), Inches(11.4), Inches(0.52),
         title, 19, WHITE, True)
    if meta:
        text(slide, Inches(9.60), Inches(0.22), Inches(3.45), Inches(0.34),
             meta, 11.5, META, False, PP_ALIGN.RIGHT)
    if page is not None:
        text(slide, Inches(12.55), Inches(7.02), Inches(0.50), Inches(0.30),
             str(page), 10.5, SUB, False, PP_ALIGN.RIGHT)


def bullets(slide, x, y, w, h, items, size=12.5, lead_color=INK):
    """items: (앞머리, 설명) 목록. 앞머리는 굵게, 설명은 보조색."""
    # 마커는 U+2022(•)를 앞머리와 같은 run에 둔다. 서식 원본이 쓰는 U+25AA(▪)는
    # 한글 글꼴에만 있어, 뒤따르는 글자가 로마자인 줄에서는 글꼴 대체가 갈리며
    # 마커가 통째로 사라진다(수식으로 시작하는 줄에서 실제로 발생했다).
    paras = []
    for lead, rest in items:
        run = [("• " + lead, True, lead_color)]
        if rest:
            run.append(("  " + rest, False, SUB))
        paras.append(run)
    return text(slide, x, y, w, h, paras, size, spacing=1.22)


def panel(slide, x, y, w, h, heading, accent, items, size=12.5, fill=PANEL):
    rect(slide, x, y, w, h, fill)
    text(slide, x + Inches(0.18), y + Inches(0.13), w - Inches(0.36),
         Inches(0.32), heading, 14, accent, True)
    if items:
        bullets(slide, x + Inches(0.18), y + Inches(0.58), w - Inches(0.36),
                h - Inches(0.72), items, size)


def dark(slide, x, y, w, h, heading, paras, size=13.5):
    rect(slide, x, y, w, h, INK)
    # 제목이 없으면 제목 자리를 비우지 않는다. 비우면 얕은 띠에서 본문이
    # 도형 아래로 밀려 슬라이드 밖으로 나간다.
    body_y = Inches(0.68) if heading else Inches(0.22)
    if heading:
        text(slide, x + Inches(0.28), y + Inches(0.20), w - Inches(0.56),
             Inches(0.36), heading, 14, WHITE, True)
    text(slide, x + Inches(0.28), y + body_y, w - Inches(0.56),
         h - body_y - Inches(0.18), paras, size, WHITE, spacing=1.30,
         anchor=MSO_ANCHOR.MIDDLE if not heading else MSO_ANCHOR.TOP)


def figure(slide, name, x, y, max_w, max_h, caption=None):
    path = os.path.join(FIG, name)
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(path, int(x + (max_w - w) / 2), y, w, h)
    if caption:
        text(slide, x, y + h + Inches(0.10), max_w, Inches(0.30),
             caption, 10.5, SUB, False, PP_ALIGN.CENTER)


def table(slide, x, y, w, rows, col_w, highlight=(), size=11.0, row_h=Inches(0.30)):
    """rows[0]은 머리행. highlight는 강조할 본문 행 인덱스(0부터)."""
    n_r, n_c = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_r, n_c, x, y, w, row_h * n_r)
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False
    total = sum(col_w)
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Emu(int(w * cw / total))
    for i, row in enumerate(rows):
        tbl.rows[i].height = row_h
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = 0
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = INK
            elif (i - 1) in highlight:
                cell.fill.fore_color.rgb = RGBColor(0xE3, 0xEE, 0xE4)
            else:
                cell.fill.fore_color.rgb = WHITE if i % 2 else PANEL
            tf = cell.text_frame
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = (i == 0) or ((i - 1) in highlight)
            r.font.color.rgb = WHITE if i == 0 else INK
    return tbl


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ---------------------------------------------------------------- 슬라이드
def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    MET = "KRRI · 2026"

    # 1 --------------------------------------------------------- 표지
    s = blank(prs)
    rect(s, 0, 0, W, Inches(2.05), INK)
    with Image.open(LOGO) as im:
        lw, lh = im.size
    logo_w = Inches(3.00)
    rect(s, Inches(0.30), Inches(0.30), logo_w + Inches(0.16),
         logo_w * lh / lw + Inches(0.16), WHITE)
    s.shapes.add_picture(LOGO, Inches(0.38), Inches(0.38), logo_w, logo_w * lh / lw)
    text(s, MARGIN, Inches(2.55), Inches(12.4), Inches(1.5),
         "철도 자기 경로 검출의 학습형 시간 정련은\n이동평균 기준선을 능가하는가", 30, INK, True, spacing=1.22)
    rect(s, MARGIN, Inches(3.95), Inches(1.5), Inches(0.05), RED)
    text(s, MARGIN, Inches(4.18), Inches(12.4), Inches(0.7),
         "Does Learned Temporal Refinement Outperform a Moving-Average Baseline "
         "in Railway Ego-Path Detection?", 14, SUB, False, spacing=1.25)
    text(s, MARGIN, Inches(5.02), Inches(12.4), Inches(0.45),
         "분기부 평가 데이터셋 구축과 통제 비교", 15, TEAL, True)
    text(s, MARGIN, Inches(5.95), Inches(12.4), Inches(0.9),
         [[("김백현", True, INK), ("  ·  ", False, SUB), ("황현철", True, INK),
           ("†", False, RED)],
          [("한국철도기술연구원  Korea Railroad Research Institute", False, SUB)]],
         14, spacing=1.25)

    # 2 --------------------------------------------------------- 연구 배경
    s = blank(prs)
    header(s, "연구 배경 — 자기 경로 검출이 가장 어려운 지점", MET, 2)
    panel(s, MARGIN, BODY_TOP, Inches(6.35), Inches(2.55),
          "분기부의 구조적 난점", RED,
          [("갈라지는 유일한 지점", "진로 판단이 가장 중요한 곳"),
           ("통과 직후 근거가 사라짐", "진로를 결정하는 텅레일이 화면 아래로 이탈"),
           ("현재 프레임만으로는 불능", "어느 갈래를 택했는지 확정할 수 없음")])
    panel(s, Inches(6.98), BODY_TOP, Inches(6.35), Inches(2.55),
          "선행 연구가 남긴 숙제", TEAL,
          [("Laurent[1]", "이 구간을 “불확실 기간”으로 규정"),
           ("제안", "순환 신경망(RNN) 결합으로 보완"),
           ("미해결", "“포괄적 시계열 데이터셋이 없어” 검증 못 함")])
    dark(s, MARGIN, Inches(4.35), CONTENT_W, Inches(2.35),
         "검증에 필요한 두 가지",
         [[("① 시계열 평가 데이터 ", True, LIME),
           ("— 분기기 통과 전후를 담은 시퀀스가 있어야 한다.", False, WHITE)],
          [("② 적절한 비교 기준 ", True, SKY),
           ("— 흔히 간과되는 쪽이다. 학습형 시간 모듈은 통상 단일 프레임 모델과만 "
            "비교되고, 학습이 필요 없는 시간 필터와는 비교되지 않는다.", False, WHITE)],
          [("시간적 안정성은 이동평균만으로도 얻을 수 있다. 학습형 모듈이 그 기준선을 "
            "넘지 못한다면 추가 파라미터와 학습 비용은 정당화되지 않는다.", True, LIME)]], 13.5)

    # 3 --------------------------------------------------------- 기여
    s = blank(prs)
    header(s, "연구 질문과 기여", MET, 3)
    dark(s, MARGIN, BODY_TOP, CONTENT_W, Inches(1.05), "연구 질문",
         [[("학습형 시간 정련은 학습이 필요 없는 이동평균 기준선을 능가하는가?", True, LIME)]], 16)
    panel(s, MARGIN, Inches(2.80), Inches(6.35), Inches(1.90),
          "데이터", GREEN,
          [("분기부 평가 데이터셋", "276 이벤트 · 47개 원본 영상 · 18개 지역"),
           ("텅레일 식별 가능성의 정량 기준", "실효 해상도로 데이터 품질 관리")])
    panel(s, Inches(6.98), Inches(2.80), Inches(6.35), Inches(1.90),
          "규명", TEAL,
          [("시퀀스 간격 정합의 효과", "안정성 2.6배, 그러나 정확도는 불변"),
           ("인도메인 지표의 한계", "이 차이는 학습 도메인 지표로 드러나지 않음")])
    panel(s, MARGIN, Inches(4.95), CONTENT_W, Inches(1.90),
          "핵심 비교와 부정 결과", RED,
          [("학습형 정련 대 이동평균", "거의 같은 정확도 비용에서 EMA가 약 6배의 안정성 개선"),
           ("다중 가설 정식화의 사전 검증", "가설이 0.4 px로 붕괴 — 상한조차 단일 경로를 넘지 못함"),
           ("방법론적 제안", "학습형 시간 모듈은 이동평균 기준선과 비교되어야 한다")])

    # 4 --------------------------------------------------------- 관련 연구
    s = blank(prs)
    header(s, "관련 연구", MET, 4)
    panel(s, MARGIN, BODY_TOP, Inches(6.35), Inches(2.45),
          "자기 경로 검출", TEAL,
          [("Laurent[1]", "종단간 딥러닝으로 정식화, 백본×헤드 조합 비교"),
           ("회귀 방식", "앵커별 x좌표와 경로 상한을 직접 예측 — 디코더 없이 낮은 지연"),
           ("헤지 특성", "모호할 때 두 궤도의 중간을 예측")])
    panel(s, Inches(6.98), BODY_TOP, Inches(6.35), Inches(2.45),
          "시간적 확장", GREEN,
          [("특징맵 수준 융합", "영상 근거를 함께 사용"),
           ("출력 수준 정련 ← 본 연구 대상", "프레임별 네트워크 동결, GRU가 최근 출력 벡터를 읽어 잔차 예측"),
           ("영행렬 초기화", "학습 전에는 프레임별 예측과 동일 출력")])
    panel(s, MARGIN, Inches(4.25), Inches(6.35), Inches(2.45),
          "비학습 대안", RED,
          [("인과적 이동평균 · 지수이동평균", "추가 학습 없이 프레임 간 흔들림 감소"),
           ("대가는 지연(lag)", "대상이 실제로 변하는 순간 반응이 늦음"),
           ("본 논문", "이 고전적 대안을 학습형 정련부와 동일 지표로 비교")])
    panel(s, Inches(6.98), Inches(4.25), Inches(6.35), Inches(2.45),
          "철도 인지 데이터셋", INK,
          [("RailSem19[2]", "주행 시점 정지 영상 약 8,500장"),
           ("OSDaR23[3]", "다중 센서 연속 주행 시퀀스"),
           ("공백", "두 데이터셋 모두 분기부 통과 전후를 집중적으로 담고 있지 않음")])

    # 5 --------------------------------------------------------- 데이터셋 구성
    s = blank(prs)
    header(s, "분기부 평가 데이터셋 — 수집과 구성", MET, 5)
    panel(s, MARGIN, BODY_TOP, Inches(6.35), Inches(2.70),
          "이벤트 정의", TEAL,
          [("기관사 시점(cab view) 영상", "분기기 통과 장면을 이벤트 단위로 절단"),
           ("20초 창 · 4 fps · 81 프레임", "프레임 간격 0.25초, 폭 1280 px"),
           ("분기(facing, split)만 유지", "합류(trailing)·고정 카메라·모형철도 제외")])
    panel(s, Inches(6.98), BODY_TOP, Inches(6.35), Inches(2.70),
          "규모와 분포", GREEN,
          [("276 이벤트", "47개 원본 영상 · 18개 지역"),
           ("지역 분포", "영국 77 · 일본 34 · 슬로베니아 25 · 독일 22 · 한국 19 · 스위스 17 등"),
           ("평가 표본", "층화 표본 30 이벤트 / 2,422 프레임")])
    panel(s, MARGIN, Inches(4.50), CONTENT_W, Inches(2.20),
          "라벨과 해석상의 주의", RED,
          [("라벨 생성", "SAM 2[4] 비디오 마스크 전파 + 학습된 세그멘테이션 모델 병용, 육안 검수"),
           ("계열 분리", "라벨 생성 모델은 세그멘테이션 계열 — 비교 대상인 회귀 계열 어느 쪽에도 유리하지 않음"),
           ("절대 IoU의 의미", "자동 라벨과의 일치도로 해석하며, 본 논문의 결론은 모델 간 상대 비교에 근거")])

    # 6 --------------------------------------------------------- 실효 해상도
    s = blank(prs)
    header(s, "텅레일 식별 가능성 기준 — 명목이 아니라 실효 해상도", MET, 6)
    panel(s, MARGIN, BODY_TOP, Inches(4.55), Inches(2.45),
          "문제", RED,
          [("모두 같은 1280×720", "해상도 표기로는 판별 불가"),
           ("업스케일 영상", "화소 수는 같지만 분기기 세부가 없음"),
           ("판별 대상", "텅레일 밀착 여부")])
    figure(s, "fig_resolution.png", Inches(5.15), Inches(1.62), Inches(8.18),
           Inches(2.20), "그림 1. 좌: 재추출 전(≤720p 원본에서 절단) · 우: 동일 구간을 1080p에서 재추출")
    rect(s, MARGIN, Inches(4.25), Inches(6.35), Inches(2.45), PANEL)
    text(s, MARGIN + Inches(0.18), Inches(4.38), Inches(6.0), Inches(0.32),
         "정량 기준", 14, TEAL, True)
    rect(s, MARGIN + Inches(0.18), Inches(4.78), Inches(5.99), Inches(0.52), WHITE)
    text(s, MARGIN + Inches(0.18), Inches(4.92), Inches(5.99), Inches(0.32),
         "hi = mean | I − up(down(I, 2), 2) |", 14, INK, True, PP_ALIGN.CENTER)
    bullets(s, MARGIN + Inches(0.18), Inches(5.45), Inches(5.99), Inches(1.15),
            [("대상 영역", "중앙–하단(세로 50–95%, 가로 22–78%)"),
             ("의미", "화소율 절반을 넘는 고주파 성분의 양 — 업스케일본은 0에 가까움")], 12.0)
    dark(s, Inches(6.98), Inches(4.25), Inches(6.35), Inches(2.45),
         "육안 검증으로 확정한 임계값",
         [[("hi = 2.63", True, SKY), ("  침목이 뭉개져 텅레일 판별 불가", False, WHITE)],
          [("hi = 3.00", True, LIME), ("  레일 윤곽과 침목이 분리되어 판별 가능", False, WHITE)],
          [("따라서 채택 기준은 hi ≥ 3.0", True, LIME)]], 14)

    # 7 --------------------------------------------------------- 품질 관리 결과
    s = blank(prs)
    header(s, "데이터 품질 관리 — 새 영상 대신 같은 구간의 재추출", MET, 7)
    table(s, MARGIN, BODY_TOP, Inches(6.10),
          [["단계", "이벤트 수"],
           ["초기 수집분", "318"],
           ["실효 해상도 기준 미달 99개(31%) 제외", "219"],
           ["동일 구간 고해상도 재추출로 복구", "275"],
           ["육안 확인한 신규 수집분 반영", "276"]],
          [4.4, 1.7], highlight=(3,), row_h=Inches(0.44))
    panel(s, MARGIN, Inches(4.10), Inches(6.10), Inches(2.60),
          "근본 원인은 수집 설정이었다", RED,
          [("원본 32편 중 29편이 1080p 이상", "4K 10편 — 그런데 720p 이하로 내려받았다"),
           ("따라서 새 영상을 찾지 않았다", "동일 구간을 고해상도로 다시 잘라 56개 복구"),
           ("큐레이션 보존", "원본 시각·창 정보를 그대로 써서 분기/합류 분류가 유지됨")])
    dark(s, Inches(6.75), BODY_TOP, Inches(6.58), Inches(5.20), "결과",
         [[("최종 데이터셋 전체가 기준을 통과한다.", True, LIME)],
          [("재추출본의 실효 해상도 중앙값은 ", False, WHITE), ("4.89", True, SKY),
           ("로, 삭제 대상이었던 값(2점대)의 약 두 배다.", False, WHITE)],
          [("명목 해상도만 확인했다면 99개의 판별 불가 이벤트가 평가셋에 남고, "
            "그중 상당수는 재추출로 되살릴 수 있다는 사실도 발견되지 않았을 것이다.", False, META)]], 14)

    # 8 --------------------------------------------------------- 평가 구간
    s = blank(prs)
    header(s, "평가 프로토콜 — 구간 정의", MET, 8)
    table(s, MARGIN, BODY_TOP, Inches(6.35),
          [["구간", "범위", "의미"],
           ["접근", "−6 ~ 0 s", "분기부가 전방에 보이는 구간"],
           ["직후", "0 ~ 2.5 s", "통과 직후, 경로가 실제로 변하는 구간"],
           ["불확실", "2.5 ~ 6 s", "분기부가 시야를 벗어난 뒤, 진로 확인 전"],
           ["후기", "6 ~ 10 s", "진로가 확정된 이후"]],
          [1.3, 1.5, 4.2], highlight=(2,), row_h=Inches(0.44))
    panel(s, MARGIN, Inches(4.00), Inches(6.35), Inches(2.70),
          "왜 초 단위로 규정하는가", RED,
          [("구간 경계는 표본화율로 환산", "각 이벤트의 4 fps로 프레임 인덱스에 대응"),
           ("혼동의 대가", "프레임 단위와 초 단위를 혼동하면 구간 폭이 4배로 어긋난다"),
           ("따라서", "본 프로토콜은 모든 경계를 초로 규정한다")])
    figure(s, "fig_zones.png", Inches(6.95), Inches(1.60), Inches(6.38),
           Inches(2.05), "그림 2. 이벤트 구조와 평가 구간 정의")
    dark(s, Inches(6.95), Inches(4.00), Inches(6.38), Inches(2.70),
         "왜 불확실 구간이 관심 대상인가",
         [[("네 구간 중 시간 정보가 이득을 줄 수 있는 곳은 여기뿐이다.", True, LIME)],
          [("접근 구간에는 분기부가 화면에 있어 현재 프레임만으로 판단이 되고, "
            "후기 구간에는 궤도 정렬로 진로가 이미 확인된다.", False, WHITE)],
          [("불확실 구간은 근거가 화면에서 사라진 뒤 아직 확인되기 전 — "
            "과거 프레임 말고는 참고할 것이 없는 유일한 구간이다.", False, WHITE)]], 13.5)

    # 9 --------------------------------------------------------- 지표
    s = blank(prs)
    header(s, "평가 프로토콜 — 지표", MET, 9)
    panel(s, MARGIN, BODY_TOP, Inches(6.35), Inches(2.45),
          "IoU — 정확도", TEAL,
          [("정의", "예측한 좌·우 레일이 이루는 영역과 라벨 영역의 교집합/합집합"),
           ("한계", "라벨이 필요하며, 자동 라벨 비중이 높아 절대값은 일치도로 읽어야 함")])
    panel(s, Inches(6.98), BODY_TOP, Inches(6.35), Inches(2.45),
          "프레임 간 지터 — 안정성", GREEN,
          [("정의", "연속 두 프레임의 예측 경로를 공통 높이 구간에서 비교한 평균 |Δx| (px)"),
           ("장점", "정답 라벨이 필요 없다 — 값이 클수록 예측이 불안정"),
           ("근거", "Laurent[1]가 불확실성 지표로 제안한 “예측의 주저(hesitancy)”의 정량화")])
    dark(s, MARGIN, Inches(4.25), CONTENT_W, Inches(2.45),
         "두 지표를 함께 보아야 하는 이유",
         [[("평활은 공짜가 아니다.", True, LIME),
           (" 지터는 얼마든지 줄일 수 있지만 그 대가로 지연이 생겨 정확도를 잃는다.",
            False, WHITE)],
          [("따라서 본 논문은 안정성 개선율과 IoU 손실을 함께 제시하고, "
            "정확도 1단위를 지불하고 얻는 안정성 개선율을 ", False, WHITE),
           ("효율", True, SKY), ("로 정의해 방법 간 비교의 축으로 삼는다.", False, WHITE)]], 14)

    # 10 -------------------------------------------------------- 비교 대상
    s = blank(prs)
    header(s, "비교 대상", MET, 10)
    panel(s, MARGIN, BODY_TOP, CONTENT_W, Inches(1.00),
          "기준 모델 (base)", INK,
          [("EfficientNet-B3 백본의 회귀형 단일 프레임 검출기", "")])
    panel(s, MARGIN, Inches(2.75), Inches(6.35), Inches(3.25),
          "학습형 시간 정련 — GRU 0.46 M (base 대비 +3.2%)", RED,
          [("구조", "base 동결, 최근 5개 출력 벡터(각 129차원)를 GRU(은닉 256, 1층)가 읽어 잔차 예측"),
           ("rnn", "단일 이미지에서 크롭 창을 보간한 합성 유사 시퀀스 — 선행 구현의 기본값"),
           ("rnn-1fps", "OSDaR23(10 fps)에서 stride 10으로 표본화한 실제 시퀀스"),
           ("rnn-4fps", "동일 영상 stride 3(3.33 fps) — 평가 영상의 4 fps에 근접"),
           ("rnn-occ", "분기부 가림을 모사하는 증강으로 학습"),
           ("통제쌍", "1fps와 4fps는 표본화 간격만 다르고 나머지 설정 동일")])
    panel(s, Inches(6.98), Inches(2.75), Inches(6.35), Inches(3.25),
          "학습 불필요 시간 필터", GREEN,
          [("적용 대상", "base의 원시 예측 벡터"),
           ("좌표계 정렬", "크롭 창이 프레임마다 이동하므로 과거 예측을 현재 좌표계로 재투영 후 결합"),
           ("boxcar-K", "최근 K개 예측의 단순 평균"),
           ("EMA-α", "out[t] = α·x[t] + (1−α)·out[t−1],  out[0] = x[0]"),
           ("파라미터", "없음 — 후처리 한 줄")])
    dark(s, MARGIN, Inches(6.20), CONTENT_W, Inches(0.85), "",
         [[("두 계열은 같은 정보(과거 예측)만 본다. ", False, WHITE),
           ("차이는 그 정보를 학습으로 쓰느냐, 고정된 규칙으로 쓰느냐뿐이다.", True, LIME)]], 13.5)

    # 11 -------------------------------------------------------- 표 1 지터
    s = blank(prs)
    header(s, "결과 ① 구간별 프레임 간 지터 (px, 괄호는 base 대비 개선율)", MET, 11)
    table(s, MARGIN, BODY_TOP, CONTENT_W,
          [["모델", "접근 −6~0 s", "직후 0~2.5 s", "불확실 2.5~6 s", "후기 6~10 s"],
           ["base", "2.96", "2.45", "3.22", "3.18"],
           ["rnn", "2.96 (−0.2%)", "2.46 (−0.3%)", "3.07 (+4.5%)", "3.14 (+1.3%)"],
           ["rnn-occ", "4.15 (−40.6%)", "2.82 (−15.1%)", "3.78 (−17.6%)", "4.08 (−28.3%)"],
           ["rnn-1fps", "3.11 (−5.3%)", "2.30 (+5.8%)", "3.03 (+5.7%)", "3.12 (+1.8%)"],
           ["rnn-4fps", "3.13 (−5.8%)", "2.41 (+1.6%)", "2.74 (+14.9%)", "3.10 (+2.4%)"],
           ["base+boxcar3", "2.54 (+14.2%)", "1.92 (+21.5%)", "2.46 (+23.6%)", "2.56 (+19.4%)"],
           ["base+EMA0.5", "2.45 (+17.3%)", "1.93 (+21.2%)", "2.44 (+24.0%)", "2.37 (+25.4%)"],
           ["base+boxcar5", "2.41 (+18.6%)", "1.73 (+29.4%)", "2.16 (+32.8%)", "2.17 (+31.7%)"]],
          [2.0, 1.6, 1.6, 1.7, 1.6], highlight=(4, 6), row_h=Inches(0.355))
    dark(s, MARGIN, Inches(5.05), CONTENT_W, Inches(1.65), "읽는 법",
         [[("학습형 정련 중 최고는 rnn-4fps로 불확실 구간 ", False, WHITE),
           ("+14.9%", True, SKY),
           (". 그러나 같은 구간에서 EMA0.5는 ", False, WHITE),
           ("+24.0%", True, LIME),
           (", 필터 계열은 전 구간에서 고르게 앞선다.", False, WHITE)],
          [("rnn-occ는 모든 구간에서 base보다 나쁘다 — 6.4절에서 원인을 다룬다.", False, META)]], 13.5)

    # 12 -------------------------------------------------------- 표 2 IoU
    s = blank(prs)
    header(s, "결과 ② 구간별 IoU (라벨 대비)", MET, 12)
    table(s, MARGIN, BODY_TOP, CONTENT_W,
          [["모델", "접근", "직후", "불확실", "후기"],
           ["base", "0.9562", "0.9598", "0.9445", "0.9481"],
           ["rnn", "0.9538", "0.9577", "0.9465", "0.9471"],
           ["rnn-occ", "0.9462", "0.9542", "0.9399", "0.9399"],
           ["rnn-1fps", "0.9378", "0.9470", "0.9361", "0.9346"],
           ["rnn-4fps", "0.9371", "0.9472", "0.9325", "0.9358"],
           ["base+boxcar3", "0.9356", "0.9372", "0.9322", "0.9314"],
           ["base+EMA0.5", "0.9365", "0.9400", "0.9349", "0.9354"],
           ["base+boxcar5", "0.9108", "0.9150", "0.9095", "0.9071"],
           ["(프레임 수 n)", "720", "300", "420", "480"]],
          [2.0, 1.6, 1.6, 1.7, 1.6], highlight=(0,), row_h=Inches(0.325))
    dark(s, MARGIN, Inches(5.20), CONTENT_W, Inches(1.50),
         "전체를 관통하는 관찰",
         [[("시간 정보의 이득은 정확도가 아니라 안정성으로만 나타난다.", True, LIME)],
          [("어떤 시간 모델도 base의 IoU를 상회하지 못한다 — 불확실 구간의 rnn이 "
            "유일한 예외이며 그 폭도 +0.0020에 그친다.", False, WHITE)]], 14)

    # 13 -------------------------------------------------------- 간격 정합
    s = blank(prs)
    header(s, "결과 ③ 시퀀스 간격 정합의 효과", MET, 13)
    panel(s, MARGIN, BODY_TOP, Inches(6.35), Inches(1.15),
          "통제쌍", INK,
          [("rnn-1fps 대 rnn-4fps", "표본화 간격만 다르고 나머지 설정은 동일")])
    table(s, MARGIN, Inches(2.95), Inches(6.35),
          [["지표", "rnn-1fps", "rnn-4fps"],
           ["불확실 구간 안정성", "+5.7%", "+14.9%"],
           ["IoU 손실", "−0.0132", "−0.0140"],
           ["학습 도메인 Test IoU", "0.98459", "0.98480"],
           ["최저 검증 손실", "0.02073", "0.02070"]],
          [3.0, 1.6, 1.7], highlight=(0,), row_h=Inches(0.42))
    dark(s, Inches(6.98), BODY_TOP, Inches(6.35), Inches(2.55),
         "간격을 맞추면 안정성만 오른다",
         [[("불확실 구간 개선이 ", False, WHITE), ("+5.7% → +14.9%", True, LIME),
           (" 로 2.6배 증가한다. 시간 정보가 실제로 필요한 구간에 효과가 집중된다.", False, WHITE)],
          [("반면 정확도는 개선되지 않는다(−0.0132 대 −0.0140으로 사실상 동일). "
            "간격 정합은 정련부가 과거 궤적을 더 잘 활용하게 하지만, "
            "그 활용이 정확도로 환산되지는 않는다.", False, WHITE)]], 13.5)
    dark(s, Inches(6.98), Inches(4.30), Inches(6.35), Inches(2.40),
         "인도메인 지표로는 보이지 않는다",
         [[("두 모델의 학습 도메인 지표는 구분되지 않는다.", True, SKY)],
          [("시간 모델의 거동은 시퀀스의 시간 구조에 의존하므로, "
            "프레임 간격이 다른 도메인에 적용할 때 인도메인 지표만으로 판단해서는 안 된다.",
            False, WHITE)]], 13.5)

    # 14 -------------------------------------------------------- 핵심 결과
    s = blank(prs)
    header(s, "결과 ④ 핵심 — 학습형 정련 대 이동평균", MET, 14)
    table(s, MARGIN, BODY_TOP, Inches(6.55),
          [["모델", "지터", "안정성", "IoU 손실", "효율"],
           ["base", "2.95", "기준", "—", "—"],
           ["rnn", "2.91", "+1.4%", "−0.0008", "(미미)"],
           ["rnn-occ", "3.71", "−25.7%", "−0.0071", "음(악화)"],
           ["rnn-1fps", "2.89", "+1.9%", "−0.0132", "144"],
           ["rnn-4fps", "2.84", "+3.6%", "−0.0140", "259"],
           ["base+boxcar3", "2.37", "+19.7%", "−0.0180", "1,094"],
           ["base+EMA0.5", "2.30", "+22.1%", "−0.0154", "1,434"],
           ["base+boxcar5", "2.12", "+28.2%", "−0.0415", "680"]],
          [1.9, 1.0, 1.2, 1.3, 1.1], highlight=(4, 6), row_h=Inches(0.355))
    figure(s, "fig_qualitative.png", Inches(7.15), Inches(1.52), Inches(6.00),
           Inches(3.45), "그림 3. (a) 불확실 구간의 한 프레임 (b) 프레임 간 흔들림의 시간 변화")
    dark(s, MARGIN, Inches(5.30), CONTENT_W, Inches(1.40),
         "0.46 M 파라미터는 정당화되지 않는다",
         [[("거의 같은 정확도 비용(−0.0140 대 −0.0154)에서 EMA는 안정성을 ", False, WHITE),
           ("약 6배", True, LIME), (" 개선한다(+3.6% 대 +22.1%). 효율로는 259 대 1,434로 ",
                                  False, WHITE), ("5.5배", True, SKY),
           (". 합성 시퀀스로 학습한 rnn은 항등 함수에 가깝게 수렴해 비용도 이득도 거의 없다.",
            False, WHITE)]], 13.5)

    # 15 -------------------------------------------------------- 부정 결과
    s = blank(prs)
    header(s, "부정 결과 — 가림 증강과 다중 가설 정식화", MET, 15)
    panel(s, MARGIN, BODY_TOP, Inches(6.35), Inches(2.75),
          "가림 증강의 재검토 (rnn-occ)", RED,
          [("모든 구간에서 최하위", "지터 −25.7%, IoU −0.0071"),
           ("원인", "증강이 모사하는 실패 모드와 실제 실패 모드의 불일치"),
           ("하단 차폐가 가르치는 것", "“경로가 가려졌다” — 그러나 실제 분기부에서 레일은 계속 보인다"),
           ("실제 불확실성", "두 갈래 중 어느 쪽인지 모르는 이산적 모호성"),
           ("방증", "정련부가 base에서 크게 이탈할수록 IoU 단조 감소 (rnn 2.36 px, rnn-occ 5.87 px)")])
    text(s, Inches(6.98), Inches(1.52), Inches(6.35), Inches(0.32),
         "다중 가설 정식화의 사전 검증 (K=2)", 14, TEAL, True)
    table(s, Inches(6.98), Inches(1.95), Inches(6.35),
          [["구간", "단일 경로", "oracle", "oracle − 단일", "가설 간 거리"],
           ["접근", "0.9562", "0.9551", "−0.0011", "0.47 px"],
           ["직후", "0.9598", "0.9601", "+0.0003", "0.43 px"],
           ["불확실", "0.9445", "0.9365", "−0.0080", "0.41 px"],
           ["후기", "0.9481", "0.9463", "−0.0018", "0.42 px"]],
          [1.1, 1.3, 1.2, 1.5, 1.3], row_h=Inches(0.40))
    dark(s, MARGIN, Inches(4.60), CONTENT_W, Inches(2.10),
         "왜 가설이 붕괴하는가",
         [[("oracle이 단일 경로를 넘는 구간이 없고(최대 +0.0003), 두 가설의 거리는 "
            "0.41–0.47 px로 사실상 같은 경로다.", True, LIME)],
          [("승자독식 학습이 가설을 특화시키려면 이미지 조건부 분포가 다봉이어야 하는데, "
            "텅레일 방향은 이미지에 관측되므로 대부분의 학습 이미지에서 정답 갈래가 결정된다. "
            "진정으로 모호한 프레임은 분기부가 시야를 벗어난 이후뿐이며 — 그 데이터는 본 논문의 "
            "평가셋과 같은 성격이라, 학습에 쓰면 평가셋이 소진된다.", False, WHITE)]], 13.5)

    # 16 -------------------------------------------------------- 결론
    s = blank(prs)
    header(s, "결론", MET, 16)
    dark(s, MARGIN, BODY_TOP, CONTENT_W, Inches(2.05), "핵심 결론",
         [[("시간 정보의 이득은 정확도가 아니라 안정성으로만 나타났다.", True, LIME)],
          [("학습형 정련부는 단순 지수이동평균에 크게 못 미친다 — 거의 동일한 정확도 "
            "비용에서 EMA가 약 6배의 안정성 개선, 효율로는 5.5배 차이.", True, SKY)],
          [("따라서 0.46 M 파라미터와 별도 학습을 요구하는 출력 수준 시간 정련은 "
            "후처리 필터 대비 정당화되지 않는다.", False, WHITE)]], 14)
    panel(s, MARGIN, Inches(3.90), Inches(6.35), Inches(2.80),
          "기여 요약", GREEN,
          [("평가 기반 마련", "분기부 276 이벤트 데이터셋 + 실효 해상도 품질 기준"),
           ("실증적 답", "Laurent[1]가 제안만 하고 검증하지 못한 시간 확장에 대한 답"),
           ("방법론적 제안", "학습형 시간 모듈은 이동평균 기준선과 비교되어야 한다"),
           ("부정 결과 보고", "가림 증강의 역효과, 다중 가설 정식화의 붕괴")])
    panel(s, Inches(6.98), Inches(3.90), Inches(6.35), Inches(2.80),
          "향후 연구", TEAL,
          [("특징맵 수준 시간 융합", "예측 벡터만이 아니라 영상 근거를 함께 활용"),
           ("적응적 필터", "예측 불확실성에 따라 평활 강도를 조절 — 흔들릴 때만 강하게"),
           ("도메인 적응 ← 우선순위 높음", "시간축 개선 폭은 IoU ±0.005인 반면 "
            "학습 도메인(0.98)과 평가 도메인(0.95)의 격차는 그보다 훨씬 크다")])

    prs.save(OUT)
    print("wrote", OUT, f"({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    build()
