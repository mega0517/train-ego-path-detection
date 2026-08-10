"""TEP-Net 통합 개발 환경 — A4 가로 12장 + 제목 배너.

세로판을 가로로 돌리면 글줄이 길어지고, 무엇보다 화면 캡처(가로세로비 1.6)가
지면 비율에 맞는다. 세로판에서는 캡처를 폭에 맞추면 세로가 남아 아래 요소를
밀어냈다. 그래서 가로판은 좌측 텍스트, 우측 화면의 2단으로 짠다.

색·서체·줄 수 추정은 make_poster_ieie_a4에서 가져온다.

  python experiments/make_poster_system_a4l.py
"""
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import make_poster_ieie_a4 as P  # noqa: E402

MM, REPO, FIG, LOGO = P.MM, P.REPO, P.FIG, P.LOGO
INK, SUB, PANEL, WHITE, META = P.INK, P.SUB, P.PANEL, P.WHITE, P.META
RED, GREEN, TEAL, LIME, SKY = P.RED, P.GREEN, P.TEAL, P.LIME, P.SKY
FONT = P.FONT

# A4 가로
W, H = Inches(297 * MM), Inches(210 * MM)
MG = Inches(12 * MM)
HEAD_H = Inches(24 * MM)
BODY_TOP = HEAD_H + Inches(7 * MM)
BODY_BOT = H - Inches(9 * MM)
FULL_W = W - 2 * MG
LCOL_W = Inches(108 * MM)                 # 좌측 텍스트 단
RCOL_X = MG + LCOL_W + Inches(8 * MM)
RCOL_W = W - MG - RCOL_X

TITLE_PT = 44          # 각 장의 제목 15.5 mm
BODY_PT = 26           # 발표 내용 9.2 mm (규정 7~10 mm)
CAP_PT = 16


def sheet(prs, no, title, accent):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    P.rect(s, 0, 0, W, HEAD_H, INK)
    P.rect(s, 0, HEAD_H - Inches(2.4 * MM), W, Inches(2.4 * MM), accent)
    P.text(s, MG, Inches(4.5 * MM), Inches(22 * MM), Inches(10 * MM),
           "%02d" % no, 28, META, True)
    P.text(s, MG + Inches(19 * MM), Inches(3 * MM), W - Inches(40 * MM),
           Inches(16 * MM), title, TITLE_PT, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    return s


def bullets(s, x, y, w, items, size=BODY_PT):
    paras, n = [], 0
    for lead, rest in items:
        run = [("• " + lead, True, INK)] if lead else []
        if rest:
            run.append((("  " if lead else "") + rest, False, SUB))
        paras.append(run)
        n += P.lines_of("• " + lead + "  " + rest, w, size)
    h = Inches(size / 72.0 * 1.52) * n + Inches(0.14) * len(items)
    P.text(s, x, y, w, h, paras, size, spacing=1.34, space=12)
    return y + h + Inches(6 * MM)


def figure(s, x, y, w, name, caption, max_h=None):
    p = os.path.join(FIG, name)
    with Image.open(p) as im:
        iw, ih = im.size
    fw, fh = w, int(w * ih / iw)
    if max_h and fh > max_h:
        fh = int(max_h)
        fw = int(fh * iw / ih)
    s.shapes.add_picture(p, int(x + (w - fw) / 2), int(y), int(fw), int(fh))
    ch = Inches(CAP_PT / 72.0 * 1.32) * P.lines_of(caption, w, CAP_PT)
    P.text(s, x, y + fh + Inches(3 * MM), w, ch, caption, CAP_PT, SUB, spacing=1.26)
    return y + fh + Inches(3 * MM) + ch + Inches(5 * MM)


def dark(s, x, y, w, title, paras, size=BODY_PT):
    n = sum(P.lines_of(str(p if isinstance(p, str) else
                       "".join(t for t, _, _ in p)), w - Inches(14 * MM), size)
            for p in paras)
    h = Inches(14 * MM) + Inches(size / 72.0 * 1.52) * n + Inches(8 * MM)
    P.rect(s, x, y, w, h, INK)
    P.text(s, x + Inches(7 * MM), y + Inches(4.5 * MM), w - Inches(14 * MM),
           Inches(9 * MM), title, 26, WHITE, True)
    P.text(s, x + Inches(7 * MM), y + Inches(15 * MM), w - Inches(14 * MM),
           h - Inches(19 * MM), paras, size, WHITE, spacing=1.36, space=8)
    return y + h + Inches(6 * MM)


def table(s, x, y, w, rows, col_w, size=22, rh_mm=12):
    n_r, n_c = len(rows), len(rows[0])
    rh = Inches(rh_mm * MM)
    tbl = s.shapes.add_table(n_r, n_c, int(x), int(y), int(w), int(rh * n_r)).table
    tbl.first_row = False
    tbl.horz_banding = False
    tot = sum(col_w)
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Emu(int(w * cw / tot))
    for i, row in enumerate(rows):
        tbl.rows[i].height = rh
        for j, v in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(3 * MM)
            c.margin_top = c.margin_bottom = 0
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = INK if i == 0 else (WHITE if i % 2 else PANEL)
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(v)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = (i == 0)
            r.font.color.rgb = WHITE if i == 0 else INK
    return y + rh * n_r + Inches(6 * MM)


def build_banner():
    prs = Presentation()
    BW, BH = Inches(1000 * MM), Inches(210 * MM)
    prs.slide_width, prs.slide_height = int(BW), int(BH)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    P.rect(s, 0, 0, BW, BH, INK)
    P.rect(s, 0, BH - Inches(6 * MM), BW, Inches(6 * MM), TEAL)
    with Image.open(LOGO) as im:
        lw, lh = im.size
    logo_w = Inches(190 * MM)
    P.rect(s, Inches(30 * MM), Inches(20 * MM), logo_w + Inches(10 * MM),
           logo_w * lh / lw + Inches(10 * MM), WHITE)
    s.shapes.add_picture(LOGO, int(Inches(35 * MM)), int(Inches(25 * MM)),
                         int(logo_w), int(logo_w * lh / lw))
    P.text(s, Inches(30 * MM), Inches(64 * MM), BW - Inches(60 * MM), Inches(60 * MM),
           "철도 자기 경로 검출을 위한\nTEP-Net 통합 개발 환경", 80, WHITE, True,
           spacing=1.18)
    P.text(s, Inches(30 * MM), Inches(142 * MM), BW - Inches(60 * MM), Inches(16 * MM),
           "TEP-Net Integrated Development Environment for Railway Ego-Path Detection",
           34, META)
    P.text(s, Inches(30 * MM), Inches(163 * MM), BW - Inches(60 * MM), Inches(26 * MM),
           [[("김백현", True, WHITE), ("  ·  ", False, META), ("황현철", True, WHITE),
             ("†", False, SKY), ("      한국철도기술연구원", False, META)]], 72)
    out = os.path.join(REPO, "poster_system_banner.pptx")
    prs.save(out)
    print("wrote", out, "(1000 x 210 mm 제목 배너)")


def build_sheets():
    prs = Presentation()
    prs.slide_width, prs.slide_height = int(W), int(H)
    IMG_H = Inches(112 * MM)          # 우측 화면의 최대 높이

    # 01 통합 환경
    s = sheet(prs, 1, "통합 환경", RED)
    y = bullets(s, MG, BODY_TOP, LCOL_W, [
        ("과제", "라벨링 비용과 노선이 바뀌면 무너지는 정확도"),
        ("접근", "초안 생성과 검수, 학습을 한 브라우저에서 잇는다"),
    ])
    dark(s, MG, y, LCOL_W, "핵심", [
        [("도구를 오가며 생기는 형식 변환을 없앤다.", True, LIME)]])
    figure(s, RCOL_X, BODY_TOP, RCOL_W, "screen_laurent.png",
           "화면 1. 원본 저장소 뷰어. TEP-Net 구조와 README를 앱 안에서 확인한다.",
           IMG_H)

    # 02 세 계층
    s = sheet(prs, 2, "세 계층", RED)
    y = bullets(s, MG, BODY_TOP, FULL_W, [
        ("1계층 라벨링", "초안 생성과 점 단위 보정, 검수 큐로 사람이 볼 프레임을 추린다"),
        ("2계층 학습", "베이스 모델을 골라 그 위에 RNN을 전이학습한다"),
        ("3계층 추론", "단일 프레임과 시간 모델의 결과를 나란히 내고 지연을 함께 표시한다"),
    ])
    dark(s, MG, y, FULL_W, "환경", [
        [("등록된 검출 모델 8종, NVIDIA A100 80GB 가속", False, WHITE)],
        [("서버 데이터셋 폴더 18개를 앱에서 직접 관리한다.", True, SKY)]])

    # 03 라벨링 화면
    s = sheet(prs, 3, "라벨링 화면", GREEN)
    y = bullets(s, MG, BODY_TOP, LCOL_W, [
        ("전파", "첫 프레임에 프롬프트를 찍으면 폴더 전체로 퍼진다"),
        ("보정", "좌우 레일을 점 단위로 다듬는다"),
    ])
    dark(s, MG, y, LCOL_W, "이어짐", [
        [("확정 라벨을 다음 프레임 초안으로 전파한다.", True, LIME)]])
    figure(s, RCOL_X, BODY_TOP, RCOL_W, "fig1_labeling_sam2.jpg",
           "화면 2. 곡선 보간과 굵기를 조절하며 초안을 검수한다.", IMG_H)

    # 04 자동 라벨
    s = sheet(prs, 4, "자동 라벨", GREEN)
    y = bullets(s, MG, BODY_TOP, FULL_W, [
        ("모델 자동 라벨", "도메인이 맞을 때 가장 깨끗한 초안을 낸다. 시드와 추적이 없어 드리프트가 생기지 않는다"),
        ("SAM 2 전파", "레일을 학습한 적이 없어도 지정한 궤도를 시퀀스 전반에서 추적한다"),
        ("곡선 보간", "점 세 개에서 레일 한 줄을 만든다"),
    ])
    dark(s, MG, y, FULL_W, "선택 기준", [
        [("도메인이 맞으면 모델, 바뀌면 SAM 2 전파를 쓴다.", True, LIME)],
        [("어긋난 프레임만 사람이 손본다.", False, WHITE)]])

    # 05 도메인 이동
    s = sheet(prs, 5, "도메인 이동", GREEN)
    y = bullets(s, MG, BODY_TOP, LCOL_W, [
        ("실증", "일반 철도로 학습한 모델은 노면 전차에서 궤도를 놓친다"),
    ])
    dark(s, MG, y, LCOL_W, "대안", [
        [("SAM 2 전파는 첫 프레임 지정만으로", False, WHITE)],
        [("290 프레임을 추적했다.", True, LIME)]])
    figure(s, RCOL_X, BODY_TOP, RCOL_W, "fig3_model_fail_tram.jpg",
           "화면 3. 학습 모델의 초안이 자기 궤도를 벗어나 인접 궤도와 주변 차량으로 이탈한다.",
           IMG_H)

    # 06 학습 화면
    s = sheet(prs, 6, "학습 화면", TEAL)
    y = bullets(s, MG, BODY_TOP, LCOL_W, [
        ("전이학습", "베이스 위에 RNN을 얹는다"),
        ("설정", "에폭, 배치, 학습률, 학습 GPU를 지정한다"),
    ])
    dark(s, MG, y, LCOL_W, "선택지", [
        [("베이스 동결과 미세조정, GPU 전처리, 다중 GPU", False, WHITE)]])
    figure(s, RCOL_X, BODY_TOP, RCOL_W, "screen_training.png",
           "화면 4. 학습 곡선과 출력 모델명이 같은 화면에 남는다.", IMG_H)

    # 07 추론 화면
    s = sheet(prs, 7, "추론 화면", TEAL)
    y = bullets(s, MG, BODY_TOP, LCOL_W, [
        ("3분할", "원본과 단일 프레임, 시간 모델을 한 화면에 낸다"),
    ])
    dark(s, MG, y, LCOL_W, "측정된 지연", [
        [("30장 처리 5.8초", True, SKY)],
        [("단일 프레임 32.8 ms", False, WHITE)],
        [("시간 모델 33.2 ms", False, WHITE)]])
    figure(s, RCOL_X, BODY_TOP, RCOL_W, "fig2_inference_compare.jpg",
           "화면 5. 같은 입력에 두 모델을 동시에 적용하고 초당 처리량을 함께 표시한다.",
           IMG_H)

    # 08 운영 화면
    s = sheet(prs, 8, "운영 화면", TEAL)
    y = bullets(s, MG, BODY_TOP, LCOL_W, [
        ("파일 이동", "로컬과 서버를 좌우 패널로 두고 복사한다"),
        ("처리", "업로드와 다운로드, 압축 해제를 앱 안에서 한다"),
    ])
    dark(s, MG, y, LCOL_W, "규모", [
        [("서버 데이터셋 폴더 18개가", False, WHITE)],
        [("이미지 수와 함께 보인다.", True, SKY)]])
    figure(s, RCOL_X, BODY_TOP, RCOL_W, "screen_files.png",
           "화면 6. 파일 브라우저.", IMG_H)

    # 09 분기기 갤러리
    s = sheet(prs, 9, "분기기 갤러리", TEAL)
    y = bullets(s, MG, BODY_TOP, LCOL_W, [
        ("열람", "국가와 노선, 영상 ID로 추린다"),
        ("규모", "분기기 통과 이벤트 273건"),
    ])
    dark(s, MG, y, LCOL_W, "지역 분포", [
        [("영국 76, 일본 34, 슬로베니아 25,", False, WHITE)],
        [("독일 22 등 18개 지역", True, SKY)]])
    figure(s, RCOL_X, BODY_TOP, RCOL_W, "screen_switch.png",
           "화면 7. 이벤트별 대표 프레임을 바로 확인한다.", IMG_H)

    # 10 보유 데이터
    s = sheet(prs, 10, "보유 데이터", RED)
    y = table(s, MG, BODY_TOP, FULL_W,
              [["구분", "규모", "용도"],
               ["분기기 이벤트", "273", "이벤트당 81 프레임, 진로상 전량"],
               ["OSDaR23", "1,535", "연속 주행 시퀀스, 시간 모델 학습"],
               ["크롤링 분기기", "613", "재추출로 화질 보강"],
               ["노면 전차", "290", "도메인 이동 평가, 30 프레임 분리"]],
              [1.1, 0.6, 2.0])
    dark(s, MG, y, FULL_W, "검수 진행", [
        [("검수 큐 417 프레임 중 9 완료. RailSem19 rs19_val 원본 보유", True, SKY)]])

    # 11 등록 모델
    s = sheet(prs, 11, "등록 모델", RED)
    y = table(s, MG, BODY_TOP, FULL_W,
              [["모델", "검출 헤드", "백본", "시간 모델", "용도"],
               ["twinkling-rocket-21", "segmentation", "efficientnet-b3", "있음", "자동 라벨 생성"],
               ["chromatic-laughter-5", "regression", "efficientnet-b3", "있음", "시간 모델 베이스"],
               ["brilliant-horse-15", "segmentation", "resnet18", "있음", "비교용"],
               ["logical-tree-1", "regression", "resnet18", "없음", "경량 비교"],
               ["fortuitous-goat-12", "classification", "efficientnet-b3", "없음", "헤드 비교"]],
              [1.9, 1.1, 1.2, 0.7, 1.2], size=18, rh_mm=13)
    dark(s, MG, y, FULL_W, "활용", [
        [("헤드와 백본 조합을 화면에서 골라 같은 입력에 바로 적용한다.", True, LIME)]])

    # 12 검수 큐와 정리
    s = sheet(prs, 12, "검수 큐와 정리", RED)
    y = bullets(s, MG, BODY_TOP, FULL_W, [
        ("선별", "독립 학습된 모델들에게 같은 프레임을 물어 원거리 경로가 갈리는 곳을 찾는다"),
        ("규모", "전체의 12.8퍼센트인 417 프레임, 43개 이벤트에 분포한다"),
        ("정리", "구축 비용 절감과 도메인 대응, 모델 간 비교 가능성을 확보했다"),
    ])
    y = dark(s, MG, y, FULL_W, "이어질 작업", [
        [("검수 완료 후 분기부 학습셋 재조립과 도메인 적응", True, LIME)]])
    P.text(s, MG, y + Inches(2 * MM), FULL_W, Inches(12 * MM),
           "한국철도기술연구원 주요사업 철도 특화 로봇 기반 선로점검 "
           "핵심기술 개발(PK26312A1)", 18, SUB)

    out = os.path.join(REPO, "poster_system_a4l.pptx")
    prs.save(out)
    print("wrote", out, f"(A4 가로 {len(prs.slides._sldIdLst)}장)")


if __name__ == "__main__":
    build_banner()
    build_sheets()
